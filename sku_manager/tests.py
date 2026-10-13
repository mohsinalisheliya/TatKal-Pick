from django.test import TestCase
from django.contrib.auth import get_user_model
from django.urls import reverse
from io import BytesIO
from openpyxl import load_workbook
from polars import Decimal
from unittest.mock import patch
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
import tempfile
import qrcode

from .models import (
	AppSettings,
	DispatchLog,
	JewelrySKU,
	MarketplaceListingSKU,
	MarketplaceOrder,
	Platform,
	PlatformPrice,
	StockBatch,
	StorageBox,
)
from .views import COLOR_PALETTE
from .pdf_orders import (
	LabelParseError,
	extract_labels,
	parse_amazon_label_text,
	parse_meesho_label_text,
)


class SalesHoldTests(TestCase):
	def setUp(self):
		self.item = JewelrySKU.objects.create(name='Hold Test Product')
		self.platform = Platform.objects.create(name='Test Platform')

	def test_stock_sync_excludes_batches_on_hold(self):
		StockBatch.objects.create(
			sku=self.item,
			batch_no='HOLD-1',
			quantity=4,
			sales_enabled=False,
		)
		StockBatch.objects.create(
			sku=self.item,
			batch_no='SELL-1',
			quantity=2,
		)

		self.item.sync_stock_from_batches()

		self.item.refresh_from_db()
		self.assertEqual(self.item.stock, 2)

	def test_dispatch_skips_older_batch_on_hold(self):
		held_batch = StockBatch.objects.create(
			sku=self.item,
			batch_no='HOLD-1',
			quantity=4,
			sales_enabled=False,
		)
		sellable_batch = StockBatch.objects.create(
			sku=self.item,
			batch_no='SELL-1',
			quantity=2,
		)

		response = self.client.post(
			reverse('scan_dispatch', args=[self.item.sku]),
			{'platform_id': self.platform.pk},
		)

		held_batch.refresh_from_db()
		sellable_batch.refresh_from_db()
		self.item.refresh_from_db()
		self.assertEqual(response.status_code, 302)
		self.assertEqual(held_batch.quantity, 4)
		self.assertEqual(sellable_batch.quantity, 1)
		self.assertEqual(self.item.stock, 1)
		self.assertEqual(DispatchLog.objects.filter(sku=self.item).count(), 1)

	def test_toggle_sales_updates_sellable_stock(self):
		user = get_user_model().objects.create_user(username='stock-user', password='test-password')
		batch = StockBatch.objects.create(
			sku=self.item,
			batch_no='HOLD-1',
			quantity=4,
		)
		self.client.force_login(user)

		response = self.client.post(reverse('toggle_batch_sales', args=[batch.pk]))

		batch.refresh_from_db()
		self.item.refresh_from_db()
		self.assertEqual(response.status_code, 302)
		self.assertFalse(batch.sales_enabled)
		self.assertEqual(self.item.stock, 0)

		self.client.post(reverse('toggle_batch_sales', args=[batch.pk]))

		batch.refresh_from_db()
		self.item.refresh_from_db()
		self.assertTrue(batch.sales_enabled)
		self.assertEqual(self.item.stock, 4)


class BoxManagerColorTests(TestCase):
	def setUp(self):
		user = get_user_model().objects.create_user(username='box-user', password='test-password')
		self.client.force_login(user)
		self.url = reverse('box_manager')

	def test_palette_contains_72_unique_swatches_and_marks_used_colors(self):
		used_color = COLOR_PALETTE[0]
		StorageBox.objects.create(name='Existing Box', color_tag=used_color)

		response = self.client.get(self.url)

		self.assertEqual(len(COLOR_PALETTE), 72)
		self.assertEqual(len(set(COLOR_PALETTE)), 72)
		swatches = response.context['color_swatches']
		self.assertEqual(len(swatches), 72)
		used_swatch = next(swatch for swatch in swatches if swatch['hex'] == used_color)
		self.assertTrue(used_swatch['is_taken'])
		self.assertEqual(used_swatch['used_by'], 'Existing Box')

	def test_creates_box_with_valid_custom_color(self):
		response = self.client.post(self.url, {
			'box_name': 'Custom Box',
			'color_tag': '#123456',
			'description': 'Test location',
		})

		self.assertRedirects(response, self.url)
		self.assertTrue(StorageBox.objects.filter(name='Custom Box', color_tag='#123456').exists())

	def test_rejects_invalid_color_duplicate_name_and_duplicate_color(self):
		StorageBox.objects.create(name='Existing Box', color_tag='#A1B2C3')
		invalid_submissions = [
			{'box_name': 'Invalid Color', 'color_tag': 'not-a-color'},
			{'box_name': 'existing box', 'color_tag': '#123456'},
			{'box_name': 'Duplicate Color', 'color_tag': '#A1B2C3'},
		]

		for submission in invalid_submissions:
			with self.subTest(submission=submission):
				response = self.client.post(self.url, submission)
				self.assertRedirects(response, self.url)

		self.assertEqual(StorageBox.objects.count(), 1)


class ProductDetailTests(TestCase):
	def setUp(self):
		user = get_user_model().objects.create_user(username='detail-user', password='test-password')
		self.client.force_login(user)
		self.item = JewelrySKU.objects.create(name='Detail Test Product')
		self.platform = Platform.objects.create(name='Detail Platform')
		self.url = reverse('product_detail', args=[self.item.pk])

	def test_detail_page_shows_batch_stock_and_sales_metrics(self):
		StockBatch.objects.create(
			sku=self.item,
			batch_no='SELL-DETAIL',
			quantity=2,
			purchase_price='5.00',
			selling_price='10.00',
		)
		StockBatch.objects.create(
			sku=self.item,
			batch_no='HOLD-DETAIL',
			quantity=3,
			purchase_price='6.00',
			selling_price='12.00',
			sales_enabled=False,
		)
		PlatformPrice.objects.create(sku=self.item, platform=self.platform, price='8.00')
		DispatchLog.objects.create(
			sku=self.item,
			platform=self.platform,
			platform_name=self.platform.name,
			quantity=1,
			sold_price='9.00',
			stock_after=1,
		)

		response = self.client.get(self.url)

		self.assertEqual(response.status_code, 200)
		self.assertTemplateUsed(response, 'sku_manager/product_detail.html')
		self.assertEqual(response.context['sellable_units'], 2)
		self.assertEqual(response.context['held_units'], 3)
		self.assertEqual(float(response.context['cost_value']), 10.0)
		self.assertEqual(float(response.context['retail_value']), 20.0)
		self.assertEqual(response.context['total_units_sold'], 1)
		self.assertEqual(float(response.context['total_revenue']), 9.0)
		self.assertEqual(response.context['units_30d'], 1)
		platform_sales = response.context['by_platform'][0]
		self.assertEqual(platform_sales['first_sale'], response.context['last_dispatch'].dispatched_at)
		self.assertEqual(platform_sales['last_sale'], response.context['last_dispatch'].dispatched_at)
		self.assertContains(response, 'On Hold')

	def test_detail_page_shows_marketplace_listing_sku(self):
		MarketplaceListingSKU.objects.create(
			sku=self.item,
			platform=self.platform,
			listing_sku='AMZ-DETAIL-123',
		)

		response = self.client.get(self.url)

		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'AMZ-DETAIL-123')

	def test_list_status_toggle_returns_to_safe_next_path(self):
		response = self.client.post(
			reverse('toggle_product_status', args=[self.item.pk]),
			{'next': self.url},
		)

		self.assertRedirects(response, self.url)
		self.item.refresh_from_db()
		self.assertFalse(self.item.is_listed)

	def test_list_status_toggle_rejects_external_next_path(self):
		response = self.client.post(
			reverse('toggle_product_status', args=[self.item.pk]),
			{'next': '//example.com/unsafe'},
		)

		self.assertRedirects(response, reverse('product_batches', args=[self.item.pk]))


class MarketplaceListingSKUTests(TestCase):
	def setUp(self):
		user = get_user_model().objects.create_user(username='mapping-user', password='test-password')
		self.client.force_login(user)
		self.item = JewelrySKU.objects.create(name='Mapping Test Product')
		self.platform = Platform.objects.create(name='Amazon')
		self.url = reverse('product_edit_details', args=[self.item.pk])

	def product_data(self, extra=None, **fields):
		data = {
			'category': self.item.category,
			'style': self.item.style,
			'name': self.item.name,
			'design_model': self.item.design_model,
			'color': self.item.color,
			'size': self.item.size,
			'number': self.item.number,
		}
		data.update(extra or {})
		data.update(fields)
		return data

	def test_product_form_saves_marketplace_sku_mapping(self):
		response = self.client.post(self.url, self.product_data({
			f'listing_sku_source_{self.platform.pk}': MarketplaceListingSKU.SOURCE_MANUAL,
			f'listing_sku_{self.platform.pk}': 'AMZ-MAPPING-123',
		}))

		self.assertRedirects(response, reverse('product_list'))
		mapping = MarketplaceListingSKU.objects.get(sku=self.item, platform=self.platform)
		self.assertEqual(mapping.listing_sku, 'AMZ-MAPPING-123')
		self.assertEqual(mapping.effective_sku, 'AMZ-MAPPING-123')

	def test_database_sku_mode_uses_the_current_product_sku(self):
		response = self.client.post(self.url, self.product_data({
			f'listing_sku_source_{self.platform.pk}': MarketplaceListingSKU.SOURCE_DATABASE,
		}, name='Edited Mapping Product'))

		self.assertRedirects(response, reverse('product_list'))
		self.item.refresh_from_db()
		mapping = MarketplaceListingSKU.objects.get(sku=self.item, platform=self.platform)
		self.assertEqual(mapping.source, MarketplaceListingSKU.SOURCE_DATABASE)
		self.assertEqual(mapping.listing_sku, '')
		self.assertEqual(mapping.effective_sku, self.item.sku)

		self.item.name = 'Updated Again Mapping Product'
		self.item.save()
		mapping.refresh_from_db()
		self.assertEqual(mapping.effective_sku, self.item.sku)

	def test_get_form_defaults_to_database_sku_mode(self):
		response = self.client.get(reverse('product_master'))

		self.assertEqual(response.status_code, 200)
		row = next(row for row in response.context['platform_rows'] if row['platform'] == self.platform)
		self.assertEqual(row['listing_sku_source'], MarketplaceListingSKU.SOURCE_DATABASE)
		self.assertEqual(row['listing_sku'], '')
		self.assertContains(response, 'Jewelry Category Code')
		self.assertContains(response, 'NCK - Necklaces / Chains / Chokers')
		self.assertContains(response, 'BGL - Bangles / Kadas')
		self.assertNotContains(response, 'APP - Apparel / Clothing')

	def test_edit_form_preserves_legacy_category(self):
		response = self.client.get(self.url)

		self.assertEqual(response.status_code, 200)
		self.assertEqual(response.context['selected_category'], self.item.category)
		self.assertContains(response, f'Existing category - {self.item.category}')

	def test_invalid_category_is_rejected_without_changing_product(self):
		response = self.client.post(self.url, self.product_data(
			name='Should Not Save',
			category='APP',
		))

		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Choose a valid jewelry category.')
		self.item.refresh_from_db()
		self.assertEqual(self.item.name, 'Mapping Test Product')

	def test_jewelry_category_can_be_selected_and_saved(self):
		response = self.client.post(self.url, self.product_data(
			name='Updated Jewelry Product',
			category='NCK',
		))

		self.assertRedirects(response, reverse('product_list'))
		self.item.refresh_from_db()
		self.assertEqual(self.item.category, 'NCK')

	def test_duplicate_marketplace_sku_is_rejected_without_saving_stock(self):
		other_item = JewelrySKU.objects.create(name='Other Mapping Product')
		MarketplaceListingSKU.objects.create(
			sku=other_item,
			platform=self.platform,
			listing_sku='AMZ-ALREADY-USED',
		)

		response = self.client.post(self.url, self.product_data({
			f'listing_sku_source_{self.platform.pk}': MarketplaceListingSKU.SOURCE_MANUAL,
			f'listing_sku_{self.platform.pk}': 'amz-already-used',
		},
			name='Should Not Save',
		))

		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'already mapped to another product')
		self.item.refresh_from_db()
		self.assertEqual(self.item.name, 'Mapping Test Product')
		self.assertFalse(MarketplaceListingSKU.objects.filter(sku=self.item).exists())

	def test_switching_seller_sku_to_database_uses_product_sku(self):
		MarketplaceListingSKU.objects.create(
			sku=self.item,
			platform=self.platform,
			listing_sku='AMZ-OLD-SKU',
		)

		response = self.client.post(self.url, self.product_data({
			f'listing_sku_source_{self.platform.pk}': MarketplaceListingSKU.SOURCE_DATABASE,
			f'listing_sku_{self.platform.pk}': 'AMZ-OLD-SKU',
		}))

		self.assertRedirects(response, reverse('product_list'))
		mapping = MarketplaceListingSKU.objects.get(sku=self.item, platform=self.platform)
		self.assertEqual(mapping.source, MarketplaceListingSKU.SOURCE_DATABASE)
		self.assertEqual(mapping.effective_sku, self.item.sku)

	def test_manual_mode_requires_a_seller_sku(self):
		response = self.client.post(self.url, self.product_data({
			f'listing_sku_source_{self.platform.pk}': MarketplaceListingSKU.SOURCE_MANUAL,
			f'listing_sku_{self.platform.pk}': '',
		}))

		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Enter a seller SKU for Amazon')

	def test_stock_form_does_not_change_marketplace_listing_sku(self):
		mapping = MarketplaceListingSKU.objects.create(
			sku=self.item,
			platform=self.platform,
			source=MarketplaceListingSKU.SOURCE_MANUAL,
			listing_sku='AMZ-KEEP-THIS',
		)

		response = self.client.post(reverse('stock_action', args=[self.item.pk]), {
			'batch_no': 'MAP-BATCH',
			'stock_qty': '4',
			'purchase_price': '10.00',
			'selling_price': '20.00',
		})

		self.assertRedirects(response, reverse('inventory_list'))
		mapping.refresh_from_db()
		self.assertEqual(mapping.listing_sku, 'AMZ-KEEP-THIS')
		self.assertEqual(self.item.batches.get().quantity, 4)


class MeeshoPDFOrderTests(TestCase):
	def setUp(self):
		user = get_user_model().objects.create_user(username='pdf-order-user', password='test-password')
		self.client.force_login(user)
		self.item = JewelrySKU.objects.create(name='PDF Order Product')
		self.platform = Platform.objects.create(name='Meesho')
		self.older_batch = StockBatch.objects.create(
			sku=self.item,
			batch_no='PDF-OLD',
			quantity=2,
		)
		self.newer_batch = StockBatch.objects.create(
			sku=self.item,
			batch_no='PDF-NEW',
			quantity=3,
		)
		self.parsed_label = {
			'seller_sku': 'MEESHO-PDF-123',
			'order_id': '123456789012345678_1',
			'quantity': 2,
			'customer_name': 'Sample Recipient',
			'shipping_address': '10 Sample Street, Sample City',
		}
		self.upload_url = reverse('pdf_order_scan')

	def upload_label(self):
		with patch(
			'sku_manager.views.extract_labels',
			return_value=([{**self.parsed_label, 'platform': 'Meesho'}], []),
		):
			return self.client.post(self.upload_url, {
				'label_pdf': SimpleUploadedFile(
					'meesho-label.pdf',
					b'%PDF-test',
					content_type='application/pdf',
				),
			})

	def test_parser_extracts_suborder_sku_quantity_and_recipient(self):
		text = """Customer Address
Sample Recipient
10 Sample Street
Sample City, State, 123456
If undelivered, return to:
SELLER
Product Details
SKU Size Qty Color Order No.
MEESHO-PDF-123 Free Size 2 Gold 123456789012345678_1
TAX INVOICE Original For Recipient"""

		result = parse_meesho_label_text(text)

		self.assertEqual(result, {
			'seller_sku': 'MEESHO-PDF-123',
			'order_id': '123456789012345678_1',
			'quantity': 2,
			'customer_name': 'Sample Recipient',
			'shipping_address': '10 Sample Street, Sample City, State, 123456',
		})

	def test_parser_preserves_spaces_inside_meesho_seller_sku(self):
		text = """Customer Address
Sample Recipient
10 Sample Street
Sample City, State, 123456
If undelivered, return to:
SELLER
Product Details
SKU Size Qty Color Order No.
VEL -HEART-002 Free Size 1 Gold 341168064498091392_1
TAX INVOICE Original For Recipient"""

		result = parse_meesho_label_text(text)

		self.assertEqual(result['seller_sku'], 'VEL -HEART-002')
		self.assertEqual(result['quantity'], 1)

	def test_parser_rejects_labels_without_meesho_order_fields(self):
		with self.assertRaises(LabelParseError):
			parse_meesho_label_text('Unrelated PDF text')

	def test_missing_pdf_dependency_does_not_prevent_app_import_and_reports_install_help(self):
		with patch.dict('sys.modules', {'pypdf': None}):
			with self.assertRaisesMessage(LabelParseError, 'same Python environment that runs Django'):
				from .pdf_orders import extract_meesho_label
				extract_meesho_label(SimpleUploadedFile('label.pdf', b'%PDF-test'))

	def test_upload_creates_pending_review_without_changing_stock(self):
		response = self.upload_label()

		order = MarketplaceOrder.objects.get(order_id=self.parsed_label['order_id'])
		self.assertRedirects(response, reverse('pdf_order_review', args=[order.pk]))
		self.assertEqual(order.status, MarketplaceOrder.STATUS_PENDING)
		self.assertIsNone(order.sku)
		self.assertEqual(self.older_batch.quantity, 2)
		self.assertEqual(self.newer_batch.quantity, 3)
		self.assertContains(self.client.get(response.url), 'Review before dispatch')

	def test_spaced_meesho_seller_sku_matches_listing_without_spaces(self):
		MarketplaceListingSKU.objects.create(
			sku=self.item,
			platform=self.platform,
			listing_sku='VEL-HEART-002',
		)
		self.parsed_label['seller_sku'] = 'VEL -HEART-002'

		self.upload_label()

		order = MarketplaceOrder.objects.get(order_id=self.parsed_label['order_id'])
		self.assertEqual(order.seller_sku, 'VEL -HEART-002')
		self.assertEqual(order.sku, self.item)

	def test_scan_page_links_to_separate_order_history(self):
		response = self.client.get(self.upload_url)

		self.assertEqual(response.status_code, 200)
		self.assertContains(response, reverse('pdf_order_history'))
		self.assertNotContains(response, 'Recent scanned orders')

	def test_order_history_is_paginated_and_links_to_excel_export(self):
		for number in range(26):
			MarketplaceOrder.objects.create(
				platform=self.platform,
				platform_name=self.platform.name,
				order_id=f'HISTORY-{number}',
				seller_sku=f'SELLER-{number}',
				quantity=1,
			)

		first_page = self.client.get(reverse('pdf_order_history'))
		second_page = self.client.get(reverse('pdf_order_history'), {'page': 2})

		self.assertEqual(first_page.status_code, 200)
		self.assertEqual(first_page.context['page_obj'].paginator.num_pages, 2)
		self.assertEqual(len(first_page.context['orders']), 25)
		self.assertContains(first_page, reverse('pdf_order_export'))
		self.assertContains(first_page, 'Next')
		self.assertEqual(len(second_page.context['orders']), 1)
		self.assertContains(second_page, 'Previous')

	def test_excel_export_contains_all_orders_and_customer_details_as_text(self):
		for number in range(26):
			MarketplaceOrder.objects.create(
				platform=self.platform,
				platform_name=self.platform.name,
				order_id=f'EXTRA-{number}',
				seller_sku=f'EXTRA-SKU-{number}',
				quantity=1,
			)
		order = MarketplaceOrder.objects.create(
			platform=self.platform,
			platform_name=self.platform.name,
			order_id='=HYPERLINK("https://example.com")',
			seller_sku='SELLER-EXCEL',
			sku=self.item,
			quantity=2,
			customer_name='Sample Recipient',
			customer_phone='9876543210',
			shipping_address='10 Sample Street',
		)

		response = self.client.get(reverse('pdf_order_export'))

		self.assertEqual(response.status_code, 200)
		self.assertEqual(
			response['Content-Type'],
			'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
		)
		self.assertIn('attachment; filename="scanned-orders.xlsx"', response['Content-Disposition'])
		workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=False)
		sheet = workbook['Scanned Orders']
		self.assertEqual(sheet.max_row, 28)
		self.assertEqual(sheet['B2'].value, "'=HYPERLINK(\"https://example.com\")")
		self.assertEqual(sheet['F2'].value, 'Sample Recipient')
		self.assertEqual(sheet['G2'].value, '9876543210')
		self.assertEqual(sheet['H2'].value, '10 Sample Street')
		self.assertEqual(sheet['D2'].value, self.item.sku)

	def test_confirm_deducts_fifo_stock_and_saves_seller_sku_mapping(self):
		self.upload_label()
		order = MarketplaceOrder.objects.get(order_id=self.parsed_label['order_id'])

		response = self.client.post(reverse('pdf_order_review', args=[order.pk]), {
			'sku_id': self.item.pk,
			'customer_name': self.parsed_label['customer_name'],
			'shipping_address': self.parsed_label['shipping_address'],
			'save_seller_sku': 'on',
		})

		order.refresh_from_db()
		self.older_batch.refresh_from_db()
		self.newer_batch.refresh_from_db()
		self.item.refresh_from_db()
		self.assertRedirects(response, reverse('pdf_order_review', args=[order.pk]))
		self.assertEqual(order.status, MarketplaceOrder.STATUS_DISPATCHED)
		self.assertEqual(order.sku, self.item)
		self.assertEqual(self.older_batch.quantity, 0)
		self.assertEqual(self.newer_batch.quantity, 3)
		self.assertEqual(self.item.stock, 3)
		dispatch = DispatchLog.objects.get(sku=self.item)
		self.assertEqual(dispatch.quantity, 2)
		self.assertEqual(order.dispatch_log, dispatch)
		self.assertEqual(
			MarketplaceListingSKU.objects.get(sku=self.item, platform=self.platform).listing_sku,
			'MEESHO-PDF-123',
		)

	def test_reuploading_dispatched_order_does_not_deduct_stock_twice(self):
		self.upload_label()
		order = MarketplaceOrder.objects.get(order_id=self.parsed_label['order_id'])
		self.client.post(reverse('pdf_order_review', args=[order.pk]), {
			'sku_id': self.item.pk,
			'customer_name': self.parsed_label['customer_name'],
			'shipping_address': self.parsed_label['shipping_address'],
		})

		response = self.upload_label()

		self.assertRedirects(response, reverse('pdf_order_review', args=[order.pk]))
		self.assertEqual(DispatchLog.objects.filter(sku=self.item).count(), 1)
		self.older_batch.refresh_from_db()
		self.assertEqual(self.older_batch.quantity, 0)

	def test_insufficient_stock_leaves_order_pending_and_does_not_dispatch(self):
		self.older_batch.quantity = 1
		self.older_batch.save()
		self.newer_batch.quantity = 0
		self.newer_batch.save()
		self.upload_label()
		order = MarketplaceOrder.objects.get(order_id=self.parsed_label['order_id'])

		response = self.client.post(reverse('pdf_order_review', args=[order.pk]), {
			'sku_id': self.item.pk,
			'customer_name': self.parsed_label['customer_name'],
			'shipping_address': self.parsed_label['shipping_address'],
		})

		order.refresh_from_db()
		self.older_batch.refresh_from_db()
		self.assertRedirects(response, reverse('pdf_order_review', args=[order.pk]))
		self.assertEqual(order.status, MarketplaceOrder.STATUS_PENDING)
		self.assertEqual(self.older_batch.quantity, 1)
		self.assertFalse(DispatchLog.objects.filter(sku=self.item).exists())


class MultiOrderParserTests(TestCase):
	def meesho_page(self, number, quantity=1):
		return f"""Customer Address
Person {number}
{number} Sample Street
Sample City, State, 12345{number}
If undelivered, return to:
SELLER
Product Details
SKU Size Qty Color Order No.
SKU-{number} Free Size {quantity} Gold 12345678901234567{number}_1
TAX INVOICE Original For Recipient"""

	def amazon_pages(self, number, sku='AMZ-SKU', quantity=2):
		order_id = f'40{number}-1234567-7654321'
		return [
			f"""Ship To:
Amz Buyer {number}
Flat {number}, Some Road
Surat, GUJARAT 39500{number}
Order ID: {order_id}""",
			f"""Amazon.in Packing Slip
Order ID: {order_id}
Quantity {quantity}
SKU: {sku}-{number} | ASIN: B0ABC""",
		]

	def test_meesho_pages_extract_independent_orders_and_skip_unreadable_pages(self):
		pages = [self.meesho_page(n) for n in range(1, 7)]
		with patch('sku_manager.pdf_orders._read_pdf_pages', return_value=pages):
			labels, skipped = extract_labels(SimpleUploadedFile('six.pdf', b'%PDF-test'))

		self.assertEqual(len(labels), 6)
		self.assertEqual({label['platform'] for label in labels}, {'Meesho'})
		self.assertEqual(skipped, [])

		with patch(
			'sku_manager.pdf_orders._read_pdf_pages',
			return_value=[self.meesho_page(1), 'random page', *self.amazon_pages(3)],
		):
			labels, skipped = extract_labels(SimpleUploadedFile('mixed.pdf', b'%PDF-test'))
		self.assertEqual({label['platform'] for label in labels}, {'Meesho', 'Amazon'})
		self.assertEqual(skipped, [2])

	def test_amazon_label_and_packing_slip_pages_merge(self):
		pages = self.amazon_pages(1) + self.amazon_pages(2)
		with patch('sku_manager.pdf_orders._read_pdf_pages', return_value=pages):
			labels, skipped = extract_labels(SimpleUploadedFile('amazon.pdf', b'%PDF-test'))

		self.assertEqual(len(labels), 2)
		self.assertEqual(skipped, [])
		first = labels[0]
		self.assertEqual(first['platform'], 'Amazon')
		self.assertEqual(first['order_id'], '401-1234567-7654321')
		self.assertEqual(first['seller_sku'], 'AMZ-SKU-1')
		self.assertEqual(first['quantity'], 2)
		self.assertEqual(first['customer_name'], 'Amz Buyer 1')
		self.assertEqual(
			first['shipping_address'],
			'Flat 1, Some Road, Surat, GUJARAT 395001',
		)

	def test_amazon_parser_keeps_the_complete_order_id(self):
		order_id = '407-9008698-2666731'
		parsed = parse_amazon_label_text(
			f'Amazon.in Packing Slip\nOrder ID: {order_id}\nSKU: AMZ-1'
		)

		self.assertEqual(parsed['order_id'], order_id)

	def test_amazon_parser_rejects_unrelated_text_and_extract_rejects_no_orders(self):
		with self.assertRaises(LabelParseError):
			parse_amazon_label_text('Unrelated PDF text')
		with patch('sku_manager.pdf_orders._read_pdf_pages', return_value=['hello world']):
			with self.assertRaises(LabelParseError):
				extract_labels(SimpleUploadedFile('empty.pdf', b'%PDF-test'))


class MultiOrderViewTests(TestCase):
	def setUp(self):
		user = get_user_model().objects.create_user(username='multi-order-user', password='test-password')
		self.client.force_login(user)
		self.item1 = JewelrySKU.objects.create(name='Multi Order One')
		self.item2 = JewelrySKU.objects.create(name='Multi Order Two')
		self.batch1 = StockBatch.objects.create(sku=self.item1, batch_no='MULTI-1', quantity=5)
		self.batch2 = StockBatch.objects.create(sku=self.item2, batch_no='MULTI-2', quantity=1)
		self.item1.sync_stock_from_batches()
		self.item2.sync_stock_from_batches()
		self.meesho = Platform.objects.create(name='Meesho')
		self.amazon = Platform.objects.create(name='Amazon')
		MarketplaceListingSKU.objects.create(
			sku=self.item1,
			platform=self.meesho,
			source=MarketplaceListingSKU.SOURCE_MANUAL,
			listing_sku='SKU-1',
		)

	def meesho_order(self, number, quantity=1):
		return {
			'platform': 'Meesho',
			'order_id': f'12345678901234567{number}_1',
			'seller_sku': f'SKU-{number}',
			'quantity': quantity,
			'customer_name': f'Person {number}',
			'shipping_address': f'{number} Sample Street',
		}

	def upload_orders(self, labels):
		with patch('sku_manager.views.extract_labels', return_value=(labels, [])):
			return self.client.post(reverse('pdf_order_scan'), {
				'label_pdf': SimpleUploadedFile('orders.pdf', b'%PDF-test'),
			})

	def test_multi_order_upload_creates_orders_and_does_not_deduct_stock(self):
		response = self.upload_orders([
			self.meesho_order(1),
			self.meesho_order(2),
			self.meesho_order(3),
			{
				'platform': 'Amazon',
				'order_id': '401-1234567-7654321',
				'seller_sku': 'AMZ-1',
				'quantity': 1,
				'customer_name': 'Amazon Buyer',
				'shipping_address': 'Amazon Street',
			},
		])

		self.assertEqual(MarketplaceOrder.objects.count(), 4)
		self.assertEqual(MarketplaceOrder.objects.filter(platform=self.amazon).count(), 1)
		self.assertTrue(response.url.startswith(reverse('pdf_order_batch')))
		page = self.client.get(response.url)
		self.assertContains(page, '4 orders found')
		self.assertContains(page, f'id="include_{MarketplaceOrder.objects.get(seller_sku="SKU-1").pk}" checked')
		self.assertContains(page, 'No product matched this seller SKU')
		self.assertEqual(page.context['pending_count'], 4)
		self.item1.refresh_from_db()
		self.assertEqual(self.item1.stock, 5)

	def test_amazon_order_id_is_preserved_in_database_history_and_dispatch_log(self):
		order_id = '407-9008698-2666731'
		response = self.upload_orders([
			{
				'platform': 'Amazon',
				'order_id': order_id,
				'seller_sku': 'AMZ-1',
				'quantity': 1,
				'customer_name': 'Amazon Buyer',
				'shipping_address': 'Amazon Street',
			},
			self.meesho_order(1),
		])

		self.assertTrue(response.url.startswith(reverse('pdf_order_batch')))
		order = MarketplaceOrder.objects.get(platform=self.amazon, order_id=order_id)
		self.assertEqual(order.order_id, order_id)
		self.assertContains(self.client.get(response.url), order_id)
		self.assertContains(self.client.get(reverse('pdf_order_history')), order_id)
		self.assertContains(
			self.client.get(reverse('pdf_order_review', args=[order.pk])),
			order_id,
		)

		dispatch_response = self.client.post(
			reverse('pdf_order_review', args=[order.pk]),
			{
				'sku_id': self.item1.pk,
				'customer_name': 'Amazon Buyer',
				'shipping_address': 'Amazon Street',
			},
			follow=True,
		)

		order.refresh_from_db()
		self.assertEqual(order.order_id, order_id)
		self.assertContains(dispatch_response, order_id)
		self.assertContains(self.client.get(reverse('dispatch_logs')), order_id)

	def test_batch_review_hides_checkbox_for_dispatched_orders(self):
		order = MarketplaceOrder.objects.create(
			platform=self.meesho,
			platform_name=self.meesho.name,
			order_id='123456789012345678_1',
			seller_sku='SKU-1',
			sku=self.item1,
			quantity=1,
			customer_name='Person One',
			shipping_address='1 Sample Street',
			status=MarketplaceOrder.STATUS_DISPATCHED,
		)

		response = self.client.get(
			reverse('pdf_order_batch'),
			{'ids': str(order.pk)},
		)

		self.assertEqual(response.context['pending_count'], 0)
		self.assertNotContains(response, f'name="include_{order.pk}"')
		self.assertContains(response, 'Dispatched')
		self.assertNotContains(response, 'Confirm selected orders and deduct stock')

	def test_single_order_pdf_still_redirects_to_single_review(self):
		response = self.upload_orders([self.meesho_order(1)])
		order = MarketplaceOrder.objects.get()
		self.assertRedirects(response, reverse('pdf_order_review', args=[order.pk]))

	def test_batch_dispatches_available_orders_and_leaves_insufficient_order_pending(self):
		self.upload_orders([self.meesho_order(1, quantity=2), self.meesho_order(2, quantity=3)])
		order1 = MarketplaceOrder.objects.get(seller_sku='SKU-1')
		order2 = MarketplaceOrder.objects.get(seller_sku='SKU-2')

		response = self.client.post(reverse('pdf_order_batch'), {
			'ids': f'{order1.pk},{order2.pk}',
			f'include_{order1.pk}': 'on',
			f'sku_{order1.pk}': self.item1.pk,
			f'name_{order1.pk}': 'Person 1',
			f'address_{order1.pk}': 'Address 1',
			f'include_{order2.pk}': 'on',
			f'sku_{order2.pk}': self.item2.pk,
			f'name_{order2.pk}': 'Person 2',
			f'address_{order2.pk}': 'Address 2',
		}, follow=True)

		order1.refresh_from_db()
		order2.refresh_from_db()
		self.item1.refresh_from_db()
		self.item2.refresh_from_db()
		self.assertEqual(order1.status, MarketplaceOrder.STATUS_DISPATCHED)
		self.assertEqual(order2.status, MarketplaceOrder.STATUS_PENDING)
		self.assertEqual(self.item1.stock, 3)
		self.assertEqual(self.item2.stock, 1)
		self.assertContains(response, 'Not enough sellable stock')

	def test_rescanning_same_orders_does_not_duplicate_or_rededuct(self):
		labels = [self.meesho_order(1), self.meesho_order(2)]
		self.upload_orders(labels)
		order1 = MarketplaceOrder.objects.get(seller_sku='SKU-1')
		self.client.post(reverse('pdf_order_batch'), {
			'ids': str(order1.pk),
			f'include_{order1.pk}': 'on',
			f'sku_{order1.pk}': self.item1.pk,
			f'name_{order1.pk}': 'Person 1',
			f'address_{order1.pk}': 'Address 1',
		})

		self.upload_orders(labels)

		self.assertEqual(MarketplaceOrder.objects.count(), 2)
		self.assertEqual(DispatchLog.objects.count(), 1)
		self.item1.refresh_from_db()
		self.assertEqual(self.item1.stock, 4)


class LabelPrintingTests(TestCase):
	def setUp(self):
		user = get_user_model().objects.create_user(username='label-user', password='test-password')
		self.client.force_login(user)
		self.item = JewelrySKU.objects.create(name='Label Test Product')
		self.older_batch = StockBatch.objects.create(
			sku=self.item,
			batch_no='OLD-LABEL',
			quantity=4,
		)
		self.latest_batch = StockBatch.objects.create(
			sku=self.item,
			batch_no='NEW-LABEL',
			quantity=2,
			section_name='Tray 7',
		)
		self.url = reverse('sku_print_label', args=[self.item.pk])

	def test_label_page_defaults_to_latest_batch_and_renders_copies(self):
		response = self.client.get(self.url)

		self.assertEqual(response.status_code, 200)
		self.assertEqual(response.context['batch'], self.latest_batch)
		self.assertEqual(response.context['copies'], self.latest_batch.quantity)
		self.assertEqual(response.context['rows'], [[1, 1]])
		self.assertContains(response, self.item.name)
		self.assertContains(response, self.item.sku)
		self.assertContains(response, 'NEW-LABEL')
		self.assertContains(response, 'data:image/png;base64,')
		self.assertNotContains(response, 'Visit Meesho Store')
		inventory_url = reverse('inventory_list')
		self.assertContains(response, f'href="{inventory_url}"')

	def test_label_includes_meesho_store_qr_when_url_is_configured(self):
		settings_obj = AppSettings.get_settings()
		settings_obj.meesho_store_url = 'https://www.meesho.com/store/test'
		settings_obj.save()
		payloads = []
		original_add_data = qrcode.QRCode.add_data

		def capture_payload(qr, data, *args, **kwargs):
			payloads.append(data)
			return original_add_data(qr, data, *args, **kwargs)

		with patch('sku_manager.views.qrcode.QRCode.add_data', autospec=True, side_effect=capture_payload):
			response = self.client.get(self.url)

		self.assertEqual(response.status_code, 200)
		self.assertEqual(payloads, [self.item.sku, settings_obj.meesho_store_url])
		self.assertTrue(response.context['store_qr_base64'])
		self.assertContains(response, 'Visit Meesho Store')
		self.assertContains(response, 'alt="Meesho Store QR"')

	def test_label_qr_encodes_only_the_sku(self):
		payloads = []
		original_add_data = qrcode.QRCode.add_data

		def capture_payload(qr, data, *args, **kwargs):
			payloads.append(data)
			return original_add_data(qr, data, *args, **kwargs)

		with patch('sku_manager.views.qrcode.QRCode.add_data', autospec=True, side_effect=capture_payload):
			label_response = self.client.get(self.url)

		self.assertEqual(label_response.status_code, 200)
		self.assertEqual(payloads, [self.item.sku])
		self.assertEqual(label_response.context['qr_payload'], self.item.sku)

	def test_inventory_has_print_label_action_for_each_sku(self):
		response = self.client.get(reverse('inventory_list'))

		self.assertEqual(response.status_code, 200)
		self.assertContains(response, reverse('sku_print_label', args=[self.item.pk]))
		self.assertContains(response, '>🏷️ Label</a>')

	def test_label_page_falls_back_for_unknown_batch_and_clamps_copies(self):
		response = self.client.get(self.url, {'batch': '999999', 'copies': '900'})

		self.assertEqual(response.context['batch'], self.latest_batch)
		self.assertEqual(response.context['copies'], 500)
		self.assertEqual(len(response.context['rows']), 250)
		self.assertTrue(all(row == [1, 1] for row in response.context['rows']))

	def test_save_and_print_redirects_to_new_batch_labels(self):
		response = self.client.post(reverse('stock_action', args=[self.item.pk]), {
			'batch_no': 'PRINT-ME',
			'stock_qty': '6',
			'purchase_price': '4.50',
			'selling_price': '9.00',
			'existing_box_id': '',
			'section_name': 'Tray 2',
			'then': 'label',
		})

		created_batch = StockBatch.objects.get(sku=self.item, batch_no='PRINT-ME')
		self.assertRedirects(
			response,
			f"{self.url}?batch={created_batch.pk}&copies=6",
			fetch_redirect_response=False,
		)

	@override_settings(MEDIA_ROOT=tempfile.gettempdir())
	def test_settings_upload_logo_and_label_uses_it(self):
		logo = SimpleUploadedFile('label-logo.png', b'fake-image-data', content_type='image/png')
		response = self.client.post(reverse('app_settings'), {
			'brand_name': 'TEST BRAND',
			'tagline': '',
			'currency_symbol': '$',
			'low_stock_threshold': '5',
			'support_contact': '',
			'label_logo': logo,
		})

		self.assertRedirects(response, reverse('app_settings'))
		settings_obj = AppSettings.get_settings()
		self.assertTrue(settings_obj.label_logo.name.startswith('branding/'))
		self.assertTrue(settings_obj.label_logo.name.endswith('.png'))

		label_response = self.client.get(self.url)
		self.assertEqual(label_response.status_code, 200)
		self.assertEqual(label_response.context['label_logo_url'], settings_obj.label_logo.url)
		self.assertContains(label_response, settings_obj.label_logo.url)

	@override_settings(MEDIA_ROOT=tempfile.gettempdir())
	def test_settings_can_remove_uploaded_label_logo(self):
		settings_obj = AppSettings.get_settings()
		settings_obj.label_logo.save(
			'existing-logo.png',
			SimpleUploadedFile('existing-logo.png', b'fake-image-data', content_type='image/png'),
		)

		self.client.post(reverse('app_settings'), {
			'brand_name': settings_obj.brand_name,
			'tagline': settings_obj.tagline,
			'currency_symbol': settings_obj.currency_symbol,
			'low_stock_threshold': settings_obj.low_stock_threshold,
			'support_contact': settings_obj.support_contact,
			'remove_label_logo': '1',
		})

		settings_obj.refresh_from_db()
		self.assertFalse(settings_obj.label_logo)

	def test_settings_save_store_urls_and_reject_invalid_store_urls(self):
		response = self.client.post(reverse('app_settings'), {
			'brand_name': 'TEST BRAND',
			'tagline': '',
			'currency_symbol': '$',
			'low_stock_threshold': '5',
			'support_contact': '',
			'meesho_store_url': 'https://meesho.com/store/test',
			'amazon_store_url': 'https://amazon.in/shop/test',
			'flipkart_store_url': 'https://flipkart.com/store/test',
		})

		self.assertRedirects(response, reverse('app_settings'))
		settings_obj = AppSettings.get_settings()
		self.assertEqual(settings_obj.meesho_store_url, 'https://meesho.com/store/test')
		self.assertEqual(settings_obj.amazon_store_url, 'https://amazon.in/shop/test')
		self.assertEqual(settings_obj.flipkart_store_url, 'https://flipkart.com/store/test')

		response = self.client.post(reverse('app_settings'), {
			'brand_name': 'TEST BRAND',
			'tagline': '',
			'currency_symbol': '$',
			'low_stock_threshold': '5',
			'support_contact': '',
			'meesho_store_url': 'not-a-url',
		})

		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Enter a valid http:// or https:// URL')
		settings_obj.refresh_from_db()
		self.assertEqual(settings_obj.meesho_store_url, 'https://meesho.com/store/test')

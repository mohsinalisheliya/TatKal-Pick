# sku_manager/views.py
import csv
import io
import base64
import colorsys
import re
import qrcode
from datetime import datetime, timedelta
from decimal import Decimal
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from django.core.paginator import Paginator
from django.core.exceptions import ValidationError
from django.core.validators import URLValidator
from django.shortcuts import render, redirect, get_object_or_404
from django.http import FileResponse, HttpResponse, JsonResponse
from django.contrib import messages
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from django.db import IntegrityError, transaction
from django.db.models import Q, Count, OuterRef, Subquery, Sum, Max, Min, F, DecimalField, ExpressionWrapper
from django.urls import reverse, NoReverseMatch
from django.utils.http import url_has_allowed_host_and_scheme
from django.utils import timezone
from django.db.models import Sum, F
from django.utils import timezone
from datetime import timedelta
import json
from .models import (
    AppSettings,
    DispatchLog,
    JewelrySKU,
    ManualStockReturn,
    MarketplaceLabelScan,
    MarketplaceListingSKU,
    MarketplaceOrder,
    Platform,
    PlatformPrice,
    StorageBox,
    StockBatch,
)
from .pdf_orders import LabelParseError, extract_labels

DEFAULT_PLATFORMS = ['Flipkart', 'Amazon', 'Meesho', 'Website']
JEWELRY_CATEGORIES = [
    ('BGL', 'Bangles / Kadas'),
    ('BRC', 'Bracelets'),
    ('NCK', 'Necklaces / Chains / Chokers'),
    ('EAR', 'Earrings'),
    ('SET', 'Jewelry Sets / Combos'),
    ('JHK', 'Jhumkas'),
    ('RNG', 'Rings'),
    ('PAY', 'Anklets / Payal'),
    ('TIK', 'Maang Tikka'),
]


def build_color_palette():
    colors = []
    for saturation, lightness in [(0.85, 0.55), (0.75, 0.45), (0.90, 0.65), (0.60, 0.35), (0.50, 0.75), (0.95, 0.50)]:
        for hue in range(0, 360, 30):
            red, green, blue = colorsys.hls_to_rgb(hue / 360, lightness, saturation)
            colors.append('#%02X%02X%02X' % (round(red * 255), round(green * 255), round(blue * 255)))
    return colors


COLOR_PALETTE = build_color_palette()


def get_or_seed_platforms():
    if not Platform.objects.exists():
        for name in DEFAULT_PLATFORMS:
            Platform.objects.get_or_create(name=name)
    return Platform.objects.all().order_by('id')

def get_all_boxes():
    return StorageBox.objects.all().order_by('name')

def get_next_serial():
    last_item = JewelrySKU.objects.order_by('-id').first()
    if not last_item:
        return '001'
    try:
        parts = last_item.sku.split('-')
        last_num = int(parts[-1])
        return str(last_num + 1).zfill(3)
    except (ValueError, IndexError):
        return str(last_item.id + 1).zfill(3)


#--- Dashboard View ---
@login_required(login_url='login')
def dashboard_home(request):
    """Main Dashboard with Charts and Low Stock Alerts"""
    settings_obj = AppSettings.get_settings()
    threshold = settings_obj.low_stock_threshold

    # Quick Stats
    total_products = JewelrySKU.objects.count()
    total_stock = JewelrySKU.objects.aggregate(Sum('stock'))['stock__sum'] or 0
    total_boxes = StorageBox.objects.count()

    # Low Stock Alerts
    low_stock_items = JewelrySKU.objects.filter(stock__lte=threshold, is_listed=True).order_by('stock')[:8]

    # Recent Dispatches
    recent_dispatches = DispatchLog.objects.select_related('sku').order_by('-dispatched_at')[:6]

    # Chart Data: Sales by Platform (Last 30 Days)
    thirty_days_ago = timezone.now() - timedelta(days=30)
    platform_sales = DispatchLog.objects.filter(dispatched_at__gte=thirty_days_ago) \
        .values('platform_name') \
        .annotate(total_units=Sum('quantity'), total_revenue=Sum(F('quantity') * F('sold_price'))) \
        .order_by('-total_units')

    chart_labels = [p['platform_name'] for p in platform_sales]
    chart_data = [p['total_units'] for p in platform_sales]

    return render(request, 'sku_manager/index.html', {
        'active_page': 'dashboard',
        'total_products': total_products,
        'total_stock': total_stock,
        'total_boxes': total_boxes,
        'low_stock_items': low_stock_items,
        'recent_dispatches': recent_dispatches,
        'currency': settings_obj.currency_symbol,
        'threshold': threshold,
        'chart_labels': json.dumps(chart_labels),
        'chart_data': json.dumps(chart_data),
    })


# --- Authentication Views ---
def user_login(request):
    if request.user.is_authenticated:
        return redirect('inventory_list')
    if request.method == 'POST':
        username = request.POST.get('username', '').strip()
        password = request.POST.get('password', '').strip()
        user = authenticate(request, username=username, password=password)
        if user is not None:
            login(request, user)
            messages.success(request, f"Welcome back, {user.username}!")
            next_url = request.POST.get('next') or request.GET.get('next') or 'inventory_list'
            return redirect(next_url)
        else:
            messages.error(request, "Invalid username or password. Please try again.")
    return render(request, 'sku_manager/login.html')

def user_logout(request):
    logout(request)
    messages.info(request, "You have been logged out successfully.")
    return redirect('login')

# --- PART 1: Product Master (Add Form + Listing Table) ---

@login_required(login_url='login')
def product_master(request, pk=None):
    """
    Unified Product Studio: Handles both Add New Product & Edit Product Details
    """
    is_edit = pk is not None
    product = get_object_or_404(JewelrySKU, pk=pk) if is_edit else None
    platforms = get_or_seed_platforms()

    def get_marketplace_rows(form_values=None):
        mapping_map = {
            mapping.platform_id: mapping
            for mapping in product.marketplace_listing_skus.all()
        } if product else {}
        rows = []
        for platform in platforms:
            mapping = mapping_map.get(platform.id)
            if form_values is not None:
                listing_sku = form_values.get(f'listing_sku_{platform.id}', '').strip()
                source = form_values.get(
                    f'listing_sku_source_{platform.id}',
                    MarketplaceListingSKU.SOURCE_MANUAL
                    if listing_sku
                    else MarketplaceListingSKU.SOURCE_DATABASE,
                )
            else:
                listing_sku = mapping.listing_sku if mapping else ''
                source = mapping.source if mapping else MarketplaceListingSKU.SOURCE_DATABASE
            rows.append({
                'platform': platform,
                'listing_sku': listing_sku,
                'listing_sku_source': source,
            })
        return rows

    form_errors = []
    form_values = None

    if request.method == 'POST':
        form_values = request.POST
        category = request.POST.get('category', '').strip().upper()
        style = request.POST.get('style', '').strip().upper()
        name = request.POST.get('name', '').strip().upper()
        design_model = request.POST.get('design_model', '').strip()
        color = request.POST.get('color', '').strip().upper()
        size = request.POST.get('size', '').strip().upper()
        number = request.POST.get('number', '').strip()

        is_existing_legacy_category = (
            product is not None
            and category.casefold() == product.category.casefold()
        )
        if category not in {code for code, _ in JEWELRY_CATEGORIES} and not is_existing_legacy_category:
            form_errors.append('Choose a valid jewelry category.')

        if not name:
            messages.error(request, "Product name is required.")
        else:
            if not number:
                number = product.number if is_edit else get_next_serial()
            number = number.zfill(3) if number.isdigit() and len(number) < 3 else number

            if is_edit:
                candidate = product
            else:
                candidate = JewelrySKU()

            candidate.category = (
                product.category
                if is_existing_legacy_category
                else category or (product.category if product else 'GEN')
            )
            candidate.style = style or (product.style if product else 'STD')
            candidate.name = name
            candidate.design_model = design_model
            candidate.color = color or (product.color if product else 'BLK')
            candidate.size = size or (product.size if product else 'FREE')
            candidate.number = number
            candidate_sku = candidate.generate_sku()
            if JewelrySKU.objects.filter(sku=candidate_sku).exclude(
                pk=product.pk if product else None,
            ).exists():
                form_errors.append(f"SKU '{candidate_sku}' already exists in the catalog.")

            listing_skus = {}
            for platform in platforms:
                listing_sku = request.POST.get(f'listing_sku_{platform.id}', '').strip()
                current_mapping = (
                    MarketplaceListingSKU.objects.filter(sku=product, platform=platform).first()
                    if product else None
                )
                default_source = (
                    MarketplaceListingSKU.SOURCE_MANUAL
                    if listing_sku
                    else current_mapping.source
                    if current_mapping
                    else MarketplaceListingSKU.SOURCE_DATABASE
                )
                source = request.POST.get(
                    f'listing_sku_source_{platform.id}',
                    default_source,
                ).strip()

                if source not in {choice[0] for choice in MarketplaceListingSKU.SOURCE_CHOICES}:
                    form_errors.append(f"Choose a valid seller SKU source for {platform.name}.")
                elif source == MarketplaceListingSKU.SOURCE_MANUAL and not listing_sku:
                    form_errors.append(f"Enter a seller SKU for {platform.name} or choose From database.")
                elif source == MarketplaceListingSKU.SOURCE_MANUAL and len(listing_sku) > 100:
                    form_errors.append(f"{platform.name} seller SKU must be 100 characters or fewer.")
                else:
                    effective_sku = (
                        listing_sku
                        if source == MarketplaceListingSKU.SOURCE_MANUAL
                        else candidate_sku
                    )
                    existing_mapping = MarketplaceListingSKU.objects.filter(
                        platform=platform,
                    ).exclude(sku=product).filter(
                        Q(source=MarketplaceListingSKU.SOURCE_MANUAL, listing_sku__iexact=effective_sku)
                        | Q(source=MarketplaceListingSKU.SOURCE_DATABASE, sku__sku__iexact=effective_sku)
                    )
                    if existing_mapping.exists():
                        form_errors.append(
                            f"'{effective_sku}' is already mapped to another product on {platform.name}."
                        )

                listing_skus[platform.id] = {
                    'source': source,
                    'listing_sku': listing_sku,
                }

            if not form_errors:
                if request.FILES.get('image'):
                    candidate.image = request.FILES['image']

                try:
                    with transaction.atomic():
                        candidate.save()
                        for platform in platforms:
                            mapping = listing_skus[platform.id]
                            MarketplaceListingSKU.objects.update_or_create(
                                sku=candidate,
                                platform=platform,
                                defaults={
                                    'source': mapping['source'],
                                    'listing_sku': (
                                        mapping['listing_sku']
                                        if mapping['source'] == MarketplaceListingSKU.SOURCE_MANUAL
                                        else ''
                                    ),
                                },
                            )
                except IntegrityError:
                    form_errors.append(
                        'A seller SKU mapping conflict prevented saving. '
                        'Check that each seller SKU is unique on its marketplace.'
                    )
                else:
                    if is_edit:
                        messages.success(request, f"✓ Product '{candidate.sku}' updated successfully!")
                        return redirect('product_list')
                    messages.success(request, f"✓ Product created successfully! SKU: {candidate.sku}")
                    return redirect('product_master')

    # Exclude current SKU from duplicate detection during edit
    if is_edit:
        existing_skus = list(JewelrySKU.objects.exclude(pk=pk).values_list('sku', flat=True))
    else:
        existing_skus = list(JewelrySKU.objects.values_list('sku', flat=True))

    products_count = JewelrySKU.objects.count()

    return render(request, 'sku_manager/products.html', {
        'is_edit': is_edit,
        'product': product,
        'existing_skus': existing_skus,
        'next_serial': product.number if is_edit else get_next_serial(),
        'products_count': products_count,
        'platform_rows': get_marketplace_rows(form_values),
        'form_errors': form_errors,
        'jewelry_categories': JEWELRY_CATEGORIES,
        'jewelry_category_codes': [code for code, _ in JEWELRY_CATEGORIES],
        'selected_category': (
            form_values.get('category', '').strip().upper()
            if form_values is not None
            else product.category if product else ''
        ),
        'active_page': 'products',
    })


@login_required(login_url='login')
def product_edit_details(request, pk):
    product = get_object_or_404(JewelrySKU, pk=pk)

    if request.method == 'POST':
        product.category = request.POST.get('category', product.category).strip().upper()
        product.style = request.POST.get('style', product.style).strip().upper()
        product.name = request.POST.get('name', product.name).strip().upper()
        product.design_model = request.POST.get('design_model', '').strip()
        product.color = request.POST.get('color', product.color).strip().upper()
        product.size = request.POST.get('size', product.size).strip().upper()
        product.number = request.POST.get('number', product.number).strip() or product.number

        if request.FILES.get('image'):
            product.image = request.FILES['image']

        product.save()
        messages.success(request, f"✓ Details updated for {product.sku}!")
        return redirect('product_master')

    existing_skus = list(JewelrySKU.objects.exclude(pk=pk).values_list('sku', flat=True))

    return render(request, 'sku_manager/product_edit.html', {
        'product': product,
        'existing_skus': existing_skus,
        'active_page': 'products',
    })


@login_required(login_url='login')
def product_list_view(request):
    """Product Listing with Search & Server-Side Pagination"""
    search_query = request.GET.get('q', '').strip()
    limit_raw = request.GET.get('limit', '10')

    try:
        limit = int(limit_raw)
    except (ValueError, TypeError):
        limit = 10
    limit = max(5, min(limit, 100))

    products = JewelrySKU.objects.all().order_by('-created_at')

    if search_query:
        products = products.filter(
            Q(sku__icontains=search_query) |
            Q(name__icontains=search_query) |
            Q(category__icontains=search_query) |
            Q(style__icontains=search_query) |
            Q(color__icontains=search_query) |
            Q(design_model__icontains=search_query)
        )

    total_count = products.count()
    paginator = Paginator(products, limit)
    page_number = request.GET.get('page', 1)
    page_obj = paginator.get_page(page_number)

    return render(request, 'sku_manager/product_list.html', {
        'products': page_obj,
        'page_obj': page_obj,
        'search_query': search_query,
        'limit': limit,
        'total_count': total_count,
        'active_page': 'products',
    })

@login_required(login_url='login')
def product_delete_secure(request, pk):
    """Delete a product only after confirming the current user's password."""
    product = get_object_or_404(JewelrySKU, pk=pk)

    if request.method == 'POST':
        entered_password = request.POST.get('delete_password', '').strip()
        if request.user.check_password(entered_password):
            sku_name = product.sku
            product.delete()
            messages.success(request, f"Product '{sku_name}' was permanently deleted.")
        else:
            messages.error(request, 'Access denied: incorrect password. Product was not deleted.')

    return redirect('product_list')

# --- PART 2: Inventory List (FIXES EMPTY TABLE) ---
@login_required(login_url='login')
def inventory_list(request):
    threshold = AppSettings.get_settings().low_stock_threshold

    for sku in JewelrySKU.objects.all():
        sku.sync_stock_from_batches()

    latest_active = (StockBatch.objects
                     .filter(sku=OuterRef('pk'), status='active', quantity__gt=0)
                     .order_by('-created_at').values('id')[:1])

    items = (JewelrySKU.objects.select_related('storage_box')
             .annotate(edit_batch_id=Subquery(latest_active))
             .order_by('-id'))

    out_count = items.filter(stock=0).count()
    low_count = items.filter(stock__gt=0, stock__lte=threshold).count()
    ok_count = items.filter(stock__gt=threshold).count()

    return render(request, 'sku_manager/inventory.html', {
        'items': items,
        'total_count': out_count + low_count + ok_count,
        'out_count': out_count,
        'low_count': low_count,
        'ok_count': ok_count,
        'threshold': threshold,
        'active_page': 'inventory',
    })

sku_inventory = inventory_list


def _safe_reverse(name, *args):
    try:
        return reverse(name, args=args)
    except NoReverseMatch:
        return None


@login_required(login_url='login')
def product_detail(request, pk):
    item = get_object_or_404(JewelrySKU.objects.select_related('storage_box'), pk=pk)
    item.sync_stock_from_batches()
    threshold = AppSettings.get_settings().low_stock_threshold

    batches = list(item.batches.select_related('storage_box').order_by('-created_at'))
    sellable_batches = [batch for batch in batches if batch.status == 'active' and batch.sales_enabled]
    held_batches = [batch for batch in batches if batch.status == 'active' and not batch.sales_enabled]
    sellable_units = sum(batch.quantity for batch in sellable_batches)
    held_units = sum(batch.quantity for batch in held_batches)
    cost_value = sum((batch.quantity * batch.purchase_price for batch in sellable_batches), Decimal('0'))
    retail_value = sum((batch.quantity * batch.selling_price for batch in sellable_batches), Decimal('0'))

    if sellable_units == 0:
        stock_state = 'out'
    elif sellable_units <= threshold:
        stock_state = 'low'
    else:
        stock_state = 'ok'

    cost = item.purchase_price
    mrp = item.selling_price
    price_map = {price.platform_id: price.price for price in item.platform_prices.all()}
    listing_sku_map = {
        mapping.platform_id: mapping
        for mapping in item.marketplace_listing_skus.all()
    }
    platform_rows = []
    for platform in get_or_seed_platforms():
        price = price_map.get(platform.id)
        row = {
            'platform': platform,
            'price': price,
            'listing_sku': (
                listing_sku_map[platform.id].effective_sku
                if platform.id in listing_sku_map
                else item.sku
            ),
            'listing_sku_source': (
                listing_sku_map[platform.id].source
                if platform.id in listing_sku_map
                else MarketplaceListingSKU.SOURCE_DATABASE
            ),
            'margin': None,
            'margin_pct': None,
            'discount_pct': None,
        }
        if price:
            row['margin'] = price - cost
            row['margin_pct'] = round((price - cost) / price * 100, 1)
            if mrp:
                row['discount_pct'] = round((mrp - price) / mrp * 100, 1)
        platform_rows.append(row)

    mrp_margin = mrp - cost if mrp else None
    mrp_margin_pct = round((mrp - cost) / mrp * 100, 1) if mrp else None

    line_total = ExpressionWrapper(
        F('sold_price') * F('quantity'),
        output_field=DecimalField(max_digits=14, decimal_places=2),
    )
    logs = item.dispatch_logs.all()
    totals = logs.aggregate(units=Sum('quantity'), revenue=Sum(line_total))
    total_units_sold = totals['units'] or 0
    total_revenue = totals['revenue'] or Decimal('0')

    last30 = logs.filter(dispatched_at__gte=timezone.now() - timedelta(days=30))
    units_30d = last30.aggregate(units=Sum('quantity'))['units'] or 0
    days_left = int(sellable_units / (units_30d / 30)) if units_30d and sellable_units else None

    by_platform = (logs.values('platform_name')
                   .annotate(units=Sum('quantity'), revenue=Sum(line_total),
                             last_sale=Max('dispatched_at'), first_sale=Min('dispatched_at'))
                   .order_by('-units'))
    recent_logs = logs.select_related('platform')[:10]
    last_dispatch = logs.first()

    return render(request, 'sku_manager/product_detail.html', {
        'item': item,
        'threshold': threshold,
        'batches': batches,
        'sellable_units': sellable_units,
        'held_units': held_units,
        'cost_value': cost_value,
        'retail_value': retail_value,
        'stock_state': stock_state,
        'platform_rows': platform_rows,
        'mrp_margin': mrp_margin,
        'mrp_margin_pct': mrp_margin_pct,
        'total_units_sold': total_units_sold,
        'total_revenue': total_revenue,
        'units_30d': units_30d,
        'days_left': days_left,
        'by_platform': by_platform,
        'recent_logs': recent_logs,
        'last_dispatch': last_dispatch,
        'edit_url': _safe_reverse('product_edit_details', item.pk) or _safe_reverse('product_edit', item.pk),
        'label_url': _safe_reverse('sku_print_label', item.pk),
        'qr_url': _safe_reverse('sku_qr_download', item.pk),
        'active_page': 'inventory',
    })


# --- PART 2: Add / Edit Stock ---
@login_required(login_url='login')
def stock_action(request, pk):
    item = get_object_or_404(JewelrySKU, pk=pk)
    item.sync_stock_from_batches()

    platforms = get_or_seed_platforms()
    boxes = get_all_boxes()

    edit_id = request.GET.get('edit_batch')
    selected_batch = None

    if edit_id:
        try:
            selected_batch = item.batches.filter(pk=int(edit_id)).first()
        except ValueError:
            selected_batch = None

    is_edit = selected_batch is not None

    def render_form(form_values=None):
        price_map = {pp.platform_id: pp.price for pp in item.platform_prices.all()}
        platform_rows = [
            {
                'platform': platform,
                'price': (
                    form_values.get(f'platform_price_{platform.id}', '')
                    if form_values is not None
                    else price_map.get(platform.id, '')
                ),
            }
            for platform in platforms  
        ]
        active_batches_with_stock = list(
            item.batches.filter(status='active', quantity__gt=0).order_by('-created_at', '-pk')
        )
        return render(request, 'sku_manager/stock_action.html', {
            'item': item,
            'is_edit': is_edit,
            'selected_batch': selected_batch,
            'boxes': boxes,
            'platform_rows': platform_rows,
            'all_batches_count': item.batches.count(),
            'suggested_batch': f"BAT-{datetime.now().strftime('%y%m%d-%H%M')}",
            'form_values': form_values,
            'active_batches_with_stock': active_batches_with_stock,
            'active_page': 'inventory',
        })

    if request.method == 'POST':
        batch_no = request.POST.get('batch_no', '').strip() or f"BAT-{datetime.now().strftime('%y%m%d%H%M')}"
        qty_raw = request.POST.get('stock_qty', '0').strip()
        stock_qty = int(qty_raw) if qty_raw.isdigit() else 0

        p_cost = request.POST.get('purchase_price', '').strip()
        w_cost = request.POST.get('wholesale_price', '').strip()
        s_price = request.POST.get(
            'selling_price',
            request.POST.get('reselling_price', ''),
        ).strip()
        purchase_price = Decimal(p_cost) if p_cost else Decimal('0.00')
        wholesale_price = Decimal(w_cost) if w_cost else Decimal('0.00')
        selling_price = Decimal(s_price) if s_price else Decimal('0.00')

        box_id = request.POST.get('existing_box_id', '').strip()
        storage_box = StorageBox.objects.filter(pk=box_id).first() if box_id else None

        section_name = request.POST.get('section_name', '').strip() or 'Main Compartment'

        with transaction.atomic():
            if is_edit:
                batch = selected_batch
                batch.batch_no = batch_no
                batch.quantity = stock_qty
                batch.purchase_price = purchase_price
                batch.wholesale_price = wholesale_price
                batch.selling_price = selling_price
                batch.storage_box = storage_box
                batch.section_name = section_name
                if stock_qty > 0:
                    batch.status = 'active'
                    batch.finished_at = None
                else:
                    batch.status = 'disabled'
                    batch.finished_at = timezone.now()
                batch.save()
                success_message = f"Batch '{batch_no}' updated! Units: {stock_qty}"
            else:
                new_batch = StockBatch.objects.create(
                    sku=item, batch_no=batch_no, quantity=stock_qty,
                    purchase_price=purchase_price,
                    wholesale_price=wholesale_price,
                    selling_price=selling_price,
                    storage_box=storage_box, section_name=section_name,
                    status='active' if stock_qty > 0 else 'disabled',
                    finished_at=timezone.now() if stock_qty == 0 else None,
                )
                success_message = f"Added new batch '{batch_no}' with {stock_qty} units!"

            for platform in platforms:
                p_val = request.POST.get(f'platform_price_{platform.id}', '').strip()
                if p_val:
                    PlatformPrice.objects.update_or_create(
                        sku=item, platform=platform, defaults={'price': Decimal(p_val)})
                else:
                    PlatformPrice.objects.filter(sku=item, platform=platform).delete()

            item.sync_stock_from_batches()
        messages.success(request, success_message)
        if request.POST.get('then') == 'label' and not is_edit and 'new_batch' in locals() and new_batch.quantity > 0:
            return redirect(f"{reverse('sku_print_label', args=[item.pk])}?batch={new_batch.pk}&copies={new_batch.quantity}")
        return redirect('inventory_list')

    return render_form()


@login_required(login_url='login')
def product_batches_view(request, pk):
    item = get_object_or_404(JewelrySKU, pk=pk)
    item.sync_stock_from_batches()

    batches = item.batches.select_related('storage_box').order_by('-created_at')
    active_units = sum(b.quantity for b in batches if b.status == 'active' and b.sales_enabled)
    held_units = sum(b.quantity for b in batches if b.status == 'active' and not b.sales_enabled)
    platforms = get_or_seed_platforms()
    platform_prices = {price.platform_id: price.price for price in item.platform_prices.all()}

    return render(request, 'sku_manager/product_batches.html', {
        'item': item,
        'product': item,
        'batches': batches,
        'active_units': active_units,
        'held_units': held_units,
        'total_batches': batches.count(),
        'boxes': get_all_boxes(),
        'platforms': platforms,
        'platform_prices': platform_prices,
        'suggested_batch': f"BAT-{datetime.now().strftime('%y%m%d-%H%M')}",
        'active_page': 'inventory',
    })


@login_required(login_url='login')
def update_universal_threshold(request):
    if request.method == 'POST':
        raw_threshold = request.POST.get('low_stock_threshold', '').strip()
        if raw_threshold.isdigit() and int(raw_threshold) >= 1:
            settings_obj = AppSettings.get_settings()
            settings_obj.low_stock_threshold = int(raw_threshold)
            settings_obj.save(update_fields=['low_stock_threshold', 'updated_at'])
            messages.success(request, 'Low-stock threshold updated.')
        else:
            messages.error(request, 'Enter a whole-number threshold of at least 1.')
    return redirect('inventory_list')


@login_required(login_url='login')
def toggle_product_status(request, pk):
    product = get_object_or_404(JewelrySKU, pk=pk)
    if request.method == 'POST':
        product.is_listed = not product.is_listed
        product.save(update_fields=['is_listed', 'updated_at'])
        messages.success(request, f"Sales status updated for {product.sku}.")
    next_url = request.POST.get('next', '')
    if url_has_allowed_host_and_scheme(
        next_url,
        allowed_hosts={request.get_host()},
        require_https=request.is_secure(),
    ):
        return redirect(next_url)
    return redirect('product_batches', pk=product.pk)


@login_required(login_url='login')
def delete_batch(request, batch_id):
    batch = get_object_or_404(StockBatch.objects.select_related('sku'), pk=batch_id)
    product_id = batch.sku_id
    if request.method == 'POST':
        batch.delete()
        messages.success(request, 'Batch deleted.')
    return redirect('product_batches', pk=product_id)


@login_required(login_url='login')
def toggle_batch_sales(request, batch_id):
    batch = get_object_or_404(StockBatch.objects.select_related('sku'), pk=batch_id)
    if request.method == 'POST':
        batch.sales_enabled = not batch.sales_enabled
        batch.save(update_fields=['sales_enabled', 'updated_at'])
        batch.sku.sync_stock_from_batches()
        if batch.sales_enabled:
            messages.success(request, f"Batch '{batch.batch_no}' enabled - units can be sold again.")
        else:
            messages.warning(request, f"Batch '{batch.batch_no}' disabled - its units cannot be sold.")
    return redirect('product_batches', pk=batch.sku_id)


@login_required(login_url='login')
def last_stock_api(request, pk):
    item = get_object_or_404(JewelrySKU, pk=pk)
    item.sync_stock_from_batches()
    return JsonResponse({'stock': item.stock})


# --- Storage Units & Platforms ---
@login_required(login_url='login')
def box_manager(request):
    used_colors_map = {box.color_tag.upper(): box.name for box in StorageBox.objects.all()}

    if request.method == 'POST':
        name = request.POST.get('box_name', '').strip()
        color_tag = request.POST.get('color_tag', '#3B82F6').strip().upper()
        description = request.POST.get('description', '').strip()

        if not name:
            messages.error(request, 'Unit name is required.')
        elif not re.match(r'^#[0-9A-F]{6}$', color_tag):
            messages.error(request, 'Pick a valid color (like #3B82F6).')
        elif StorageBox.objects.filter(name__iexact=name).exists():
            messages.error(request, f"A storage unit named '{name}' already exists.")
        elif color_tag in used_colors_map:
            messages.error(request, f"That color is already used by '{used_colors_map[color_tag]}'.")
        else:
            StorageBox.objects.create(name=name, color_tag=color_tag, description=description)
            messages.success(request, f"Storage unit '{name}' created!")
        return redirect('box_manager')

    color_swatches = [
        {'hex': color, 'is_taken': color in used_colors_map, 'used_by': used_colors_map.get(color, '')}
        for color in COLOR_PALETTE
    ]
    existing_boxes = StorageBox.objects.annotate(total_skus=Count('skus')).order_by('name')

    return render(request, 'sku_manager/boxes.html', {
        'boxes': existing_boxes,
        'color_swatches': color_swatches,
        'used_colors_map': used_colors_map,
        'active_page': 'boxes',
    })

def box_delete(request, pk):
    box = get_object_or_404(StorageBox, pk=pk)
    name = box.name
    box.delete()
    messages.info(request, f"Box '{name}' deleted.")
    return redirect('box_manager')

@login_required(login_url='login')
def platform_manager(request):
    get_or_seed_platforms()
    if request.method == 'POST':
        name = request.POST.get('platform_name', '').strip()
        if name:
            Platform.objects.get_or_create(name=name)
            messages.success(request, f"Platform '{name}' saved.")
        return redirect('platform_manager')
    platforms = Platform.objects.annotate(linked_products=Count('platformprice')).order_by('-id')
    return render(request, 'sku_manager/platforms.html', {'platforms': platforms, 'active_page': 'platforms'})

def platform_delete(request, pk):
    get_object_or_404(Platform, pk=pk).delete()
    return redirect('platform_manager')

# --- Dispatch, Scanner & Label Print ---
def scan_dispatch(request, sku):
    item = get_object_or_404(JewelrySKU.objects.select_related('storage_box'), sku=sku)
    item.sync_stock_from_batches()
    platforms = get_or_seed_platforms()
    price_map = {pp.platform_id: pp.price for pp in item.platform_prices.all()}

    if request.method == 'POST':
        platform_id = request.POST.get('platform_id')
        platform = get_object_or_404(Platform, pk=platform_id)

        if item.stock <= 0:
            messages.error(request, f"Out of stock: Cannot dispatch '{item.sku}'!")
            return redirect('scan_dispatch', sku=item.sku)

        with transaction.atomic():
            active_batch = item.batches.select_for_update().filter(
                status='active', sales_enabled=True, quantity__gt=0
            ).order_by('created_at').first()
            if not active_batch:
                item.sync_stock_from_batches()
                messages.error(request, f"Out of stock: Cannot dispatch '{item.sku}'!")
                return redirect('scan_dispatch', sku=item.sku)

            active_batch.quantity -= 1
            if active_batch.quantity == 0:
                active_batch.status = 'disabled'
                active_batch.finished_at = timezone.now()
            active_batch.save()
            item.sync_stock_from_batches()

        sold_price = price_map.get(platform.id, item.selling_price)
        DispatchLog.objects.create(
            sku=item,
            platform=platform,
            platform_name=platform.name,
            quantity=1,
            sold_price=sold_price,
            stock_after=item.stock
        )
        messages.success(request, f"Dispatched 1 unit for {platform.name}! Stock left: {item.stock}")
        return redirect('scan_dispatch', sku=item.sku)

    platform_data = [{'platform': p, 'price': price_map.get(p.id, item.selling_price)} for p in platforms]
    recent_logs = item.dispatch_logs.all()[:5]

    return render(request, 'sku_manager/dispatch.html', {
        'item': item,
        'platforms': platform_data,
        'recent_logs': recent_logs,
    })

@login_required(login_url='login')
def sku_print_label(request, pk):
    item = get_object_or_404(JewelrySKU.objects.select_related('storage_box'), pk=pk)
    settings_obj = AppSettings.get_settings()
    batches = list(item.batches.order_by('-created_at', '-pk'))
    batch = None
    batch_id = request.GET.get('batch', '')
    if batch_id.isdigit():
        batch = next((entry for entry in batches if entry.pk == int(batch_id)), None)
    if batch is None:
        batch = next((entry for entry in batches if entry.status == 'active' and entry.quantity > 0), None)
        batch = batch or (batches[0] if batches else None)

    copies_raw = request.GET.get('copies', '')
    if copies_raw.isdigit():
        copies = int(copies_raw)
    else:
        copies = batch.quantity if batch and batch.quantity > 0 else 1
    copies = max(1, min(copies, 500))

    rows = [[1, 1] for _ in range(copies // 2)]
    if copies % 2:
        rows.append([1])

    qr = qrcode.QRCode(
        version=None,
        error_correction=qrcode.constants.ERROR_CORRECT_L,
        box_size=10,
        border=1,
    )
    qr.add_data(item.sku)
    qr.make(fit=True)
    img = qr.make_image(fill_color="black", back_color="white")

    buffer = io.BytesIO()
    img.save(buffer, format="PNG")
    qr_base64 = base64.b64encode(buffer.getvalue()).decode('utf-8')
    store_qr_base64 = ''
    if settings_obj.meesho_store_url:
        store_qr = qrcode.QRCode(
            version=None,
            error_correction=qrcode.constants.ERROR_CORRECT_L,
            box_size=4,
            border=1,
        )
        store_qr.add_data(settings_obj.meesho_store_url)
        store_qr.make(fit=True)
        store_img = store_qr.make_image(fill_color="black", back_color="white")
        store_buffer = io.BytesIO()
        store_img.save(store_buffer, format="PNG")
        store_qr_base64 = base64.b64encode(store_buffer.getvalue()).decode('utf-8')

    return render(request, 'sku_manager/label_print.html', {
        'item': item,
        'batch': batch,
        'batches': batches,
        'copies': copies,
        'rows': rows,
        'qr_base64': qr_base64,
        'store_qr_base64': store_qr_base64,
        'qr_payload': item.sku,
        'qr_mm': 19,
        'label_logo_url': settings_obj.label_logo.url if settings_obj.label_logo else '',
    })

def sku_qr_download(request, pk):
    item = get_object_or_404(JewelrySKU, pk=pk)
    qr_payload = item.sku
    qr = qrcode.QRCode(box_size=10, border=2)
    qr.add_data(qr_payload)
    qr.make(fit=True)
    img = qr.make_image(fill_color="black", back_color="white")
    buffer = io.BytesIO()
    img.save(buffer, format="PNG")
    response = HttpResponse(buffer.getvalue(), content_type="image/png")
    response['Content-Disposition'] = f'attachment; filename="QR_{item.sku}.png"'
    return response

@login_required(login_url='login')
def sku_scanner(request):
    return render(request, 'sku_manager/scanner.html', {'active_page': 'scanner'})

@login_required(login_url='login')
def dispatch_logs(request):
    logs = DispatchLog.objects.select_related('sku', 'marketplace_order').all()
    return render(request, 'sku_manager/dispatch_logs.html', {
        'logs': logs,
        'total_dispatches': logs.count(),
        'active_page': 'dispatch_logs',
    })

def normalize_seller_sku(value):
    return re.sub(r'\s+', '', value).casefold()


def find_product_for_seller_sku(platform, seller_sku):
    normalized_sku = normalize_seller_sku(seller_sku)
    matches = {
        product.pk: product
        for product in JewelrySKU.objects.all()
        if normalize_seller_sku(product.sku) == normalized_sku
    }
    for mapping in MarketplaceListingSKU.objects.filter(platform=platform).select_related('sku'):
        if normalize_seller_sku(mapping.effective_sku) == normalized_sku:
            matches[mapping.sku_id] = mapping.sku
    return next(iter(matches.values())) if len(matches) == 1 else None


@login_required(login_url='login')
def pdf_order_scan(request):
    if request.method == 'POST':
        uploaded_file = request.FILES.get('label_pdf')
        if not uploaded_file:
            messages.error(request, 'Choose a marketplace shipping-label PDF to scan.')
            return redirect('pdf_order_scan')

        try:
            parsed_labels, skipped_pages = extract_labels(uploaded_file)
        except LabelParseError as error:
            messages.error(request, str(error))
            return redirect('pdf_order_scan')

        orders = []
        for parsed in parsed_labels:
            platform, _ = Platform.objects.get_or_create(name=parsed['platform'])
            product = find_product_for_seller_sku(platform, parsed['seller_sku'])
            order, _ = MarketplaceOrder.objects.get_or_create(
                platform=platform,
                order_id=parsed['order_id'],
                defaults={
                    'platform_name': platform.name,
                    'seller_sku': parsed['seller_sku'],
                    'sku': product,
                    'quantity': parsed['quantity'],
                    'customer_name': parsed['customer_name'],
                    'shipping_address': parsed['shipping_address'],
                },
            )
            orders.append(order)

        uploaded_file.seek(0)
        scan = MarketplaceLabelScan()
        scan.pdf_file.save(uploaded_file.name, uploaded_file, save=True)
        scan.orders.set(orders)

        if skipped_pages:
            messages.warning(
                request,
                'Some pages could not be read: ' + ', '.join(map(str, skipped_pages)),
            )
        if len(orders) == 1:
            return redirect('pdf_order_review', pk=orders[0].pk)
        batch_url = reverse('pdf_order_batch')
        order_ids = ','.join(str(order.pk) for order in orders)
        query = f'{batch_url}?ids={order_ids}&scan_id={scan.pk}'
        if skipped_pages:
            query += '&skipped=' + ','.join(map(str, skipped_pages))
        return redirect(query)

    return render(request, 'sku_manager/pdf_order_scan.html', {
        'active_page': 'pdf_order_scan',
    })


@login_required(login_url='login')
def pdf_order_history(request):
    orders = MarketplaceOrder.objects.select_related('sku').order_by('-created_at', '-pk')
    page_obj = Paginator(orders, 25).get_page(request.GET.get('page'))
    return render(request, 'sku_manager/pdf_order_history.html', {
        'page_obj': page_obj,
        'orders': page_obj.object_list,
        'total_orders': orders.count(),
        'active_page': 'pdf_order_scan',
    })


@login_required(login_url='login')
def pdf_order_export(request):
    orders = MarketplaceOrder.objects.select_related('sku').order_by('-created_at', '-pk')
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Scanned Orders'
    sheet.append([
        'Platform',
        'Order ID',
        'Seller SKU',
        'Tatkal Pick SKU',
        'Quantity',
        'Recipient Name',
        'Customer Phone',
        'Delivery Address',
        'Status',
        'Scanned At',
        'Dispatched At',
    ])

    def safe_text(value):
        text = str(value or '')
        return f"'{text}" if text.lstrip().startswith(('=', '+', '-', '@')) else text

    for order in orders:
        sheet.append([
            safe_text(order.platform_name),
            safe_text(order.order_id),
            safe_text(order.seller_sku),
            safe_text(order.sku.sku if order.sku else ''),
            order.quantity,
            safe_text(order.customer_name),
            safe_text(order.customer_phone),
            safe_text(order.shipping_address),
            order.get_status_display(),
            timezone.localtime(order.created_at).strftime('%Y-%m-%d %H:%M:%S'),
            (
                timezone.localtime(order.dispatched_at).strftime('%Y-%m-%d %H:%M:%S')
                if order.dispatched_at else ''
            ),
        ])

    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for column_cells in sheet.columns:
        width = min(max(max(len(str(cell.value or '')) for cell in column_cells) + 2, 12), 60)
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width

    output = io.BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = 'attachment; filename="scanned-orders.xlsx"'
    return response


class MarketplaceOrderDispatchError(ValueError):
    pass


def marketplace_pdf_stamper_context(scan):
    if not scan:
        return {}

    settings_obj = AppSettings.get_settings()
    return {
        'pdf_url': reverse('pdf_order_scan_file', args=[scan.pk]),
        'pdf_stamp_url': reverse('pdf_order_stamped_file', args=[scan.pk]),
        'scan_id': scan.pk,
        'label_logo_url': settings_obj.label_logo.url if settings_obj.label_logo else '',
        'store_qr_available': bool(settings_obj.meesho_store_url),
    }


def stamp_marketplace_pdf(scan, include_logo, include_store_qr):
    from pypdf import PdfReader, PdfWriter
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas
    from PIL import Image, ImageOps

    settings_obj = AppSettings.get_settings()
    logo = None
    if include_logo and settings_obj.label_logo:
        with settings_obj.label_logo.open('rb') as logo_file:
            logo = Image.open(logo_file)
            logo.load()
            logo = ImageOps.exif_transpose(logo).convert('RGBA')
            logo.thumbnail((1200, 1200), Image.Resampling.LANCZOS)

    store_qr = None
    if include_store_qr and settings_obj.meesho_store_url:
        qr_code = qrcode.QRCode(
            version=None,
            error_correction=qrcode.constants.ERROR_CORRECT_M,
            box_size=6,
            border=2,
        )
        qr_code.add_data(settings_obj.meesho_store_url)
        qr_code.make(fit=True)
        store_qr = qr_code.make_image(fill_color='black', back_color='white').convert('RGB')

    output = io.BytesIO()
    overlay_readers = []
    writer = PdfWriter()
    with scan.pdf_file.open('rb') as source_file:
        reader = PdfReader(source_file)
        for page in reader.pages:
            if logo or store_qr:
                page_width = float(page.mediabox.width)
                page_height = float(page.mediabox.height)
                margin = min(18, page_width * 0.04, page_height * 0.04)
                overlay_buffer = io.BytesIO()
                overlay = canvas.Canvas(
                    overlay_buffer,
                    pagesize=(page_width, page_height),
                    pageCompression=1,
                )

                if logo:
                    max_logo_width = min(90, page_width * 0.28)
                    max_logo_height = min(42, page_height * 0.12)
                    image_width, image_height = logo.size
                    scale = min(
                        max_logo_width / image_width,
                        max_logo_height / image_height,
                    )
                    draw_width = image_width * scale
                    draw_height = image_height * scale
                    overlay.drawImage(
                        ImageReader(logo),
                        margin,
                        margin,
                        width=draw_width,
                        height=draw_height,
                        mask='auto',
                    )

                if store_qr:
                    qr_size = min(48, page_width * 0.16, page_height * 0.14)
                    overlay.drawImage(
                        ImageReader(store_qr),
                        page_width - margin - qr_size,
                        margin,
                        width=qr_size,
                        height=qr_size,
                    )

                overlay.showPage()
                overlay.save()
                overlay_buffer.seek(0)
                overlay_reader = PdfReader(overlay_buffer)
                overlay_readers.append((overlay_buffer, overlay_reader))
                page.merge_page(overlay_reader.pages[0])
            writer.add_page(page)

        writer.write(output)

    output.seek(0)
    return output


def dispatch_marketplace_order(order_id, sku_id, customer_name, shipping_address, save_seller_sku=False):
    with transaction.atomic():
        order = get_object_or_404(
            MarketplaceOrder.objects.select_for_update().select_related('platform'),
            pk=order_id,
        )
        if order.status != MarketplaceOrder.STATUS_PENDING:
            if order.status == MarketplaceOrder.STATUS_DISPATCHED:
                raise MarketplaceOrderDispatchError('This order was already dispatched.')
            if order.status == MarketplaceOrder.STATUS_CANCELLED:
                raise MarketplaceOrderDispatchError('This order was cancelled and cannot be dispatched.')
            raise MarketplaceOrderDispatchError('This order cannot be dispatched in its current status.')

        if not str(sku_id or '').isdigit():
            raise MarketplaceOrderDispatchError('Choose a valid Tatkal Pick product.')
        item = JewelrySKU.objects.filter(pk=sku_id).first()
        if item is None:
            raise MarketplaceOrderDispatchError('Choose a valid Tatkal Pick product.')
        customer_name = customer_name.strip()
        shipping_address = shipping_address.strip()
        if not customer_name or not shipping_address:
            raise MarketplaceOrderDispatchError('Customer name and delivery address are required.')

        batches = list(
            StockBatch.objects.select_for_update().filter(
                sku=item,
                status='active',
                sales_enabled=True,
                quantity__gt=0,
            ).order_by('created_at', 'pk')
        )
        available_stock = sum(batch.quantity for batch in batches)
        if available_stock < order.quantity:
            raise MarketplaceOrderDispatchError(
                f"Not enough sellable stock for {item.sku}: "
                f"{available_stock} available, {order.quantity} required."
            )

        remaining = order.quantity
        for batch in batches:
            dispatched = min(batch.quantity, remaining)
            batch.quantity -= dispatched
            remaining -= dispatched
            if batch.quantity == 0:
                batch.status = 'disabled'
                batch.finished_at = timezone.now()
            batch.save(update_fields=['quantity', 'status', 'finished_at', 'updated_at'])
            if remaining == 0:
                break

        item.sync_stock_from_batches()
        sold_price = PlatformPrice.objects.filter(
            sku=item,
            platform=order.platform,
        ).values_list('price', flat=True).first()
        if sold_price is None:
            sold_price = item.selling_price
        dispatch_log = DispatchLog.objects.create(
            sku=item,
            platform=order.platform,
            platform_name=order.platform.name,
            quantity=order.quantity,
            sold_price=sold_price,
            stock_after=item.stock,
        )

        if save_seller_sku:
            MarketplaceListingSKU.objects.update_or_create(
                sku=item,
                platform=order.platform,
                defaults={
                    'source': MarketplaceListingSKU.SOURCE_MANUAL,
                    'listing_sku': order.seller_sku,
                },
            )

        order.sku = item
        order.customer_name = customer_name
        order.shipping_address = shipping_address
        order.dispatch_log = dispatch_log
        order.status = MarketplaceOrder.STATUS_DISPATCHED
        order.dispatched_at = timezone.now()
        order.save(update_fields=[
            'sku',
            'customer_name',
            'shipping_address',
            'dispatch_log',
            'status',
            'dispatched_at',
        ])
    return order, item


@login_required(login_url='login')
def order_cancellation(request):
    if request.method == 'POST':
        mode = request.POST.get('mode')

        if mode == 'auto':
            order_id = request.POST.get('order_id', '').strip()
            platform_id = request.POST.get('platform_id', '').strip()
            if not order_id:
                messages.error(request, 'Please enter an Order ID.')
                return redirect('order_cancellation')

            orders = MarketplaceOrder.objects.filter(order_id=order_id)
            if platform_id:
                if not platform_id.isdigit():
                    messages.error(request, 'Choose a valid marketplace.')
                    return redirect('order_cancellation')
                orders = orders.filter(platform_id=platform_id)
            else:
                matches = list(orders.select_related('sku', 'platform')[:2])
                if len(matches) > 1:
                    messages.error(
                        request,
                        'This Order ID exists on multiple marketplaces. Choose the marketplace and try again.',
                    )
                    return redirect('order_cancellation')
                order = matches[0] if matches else None

            if platform_id:
                order = orders.select_related('sku', 'platform').first()
            if order is None:
                messages.error(
                    request,
                    f"Order '{order_id}' not found. If it was dispatched before using this software, use Manual Restock.",
                )
                return redirect('order_cancellation')

            with transaction.atomic():
                order = MarketplaceOrder.objects.select_for_update().select_related('sku').get(pk=order.pk)
                if order.status == MarketplaceOrder.STATUS_CANCELLED:
                    messages.warning(request, f"Order '{order_id}' is already cancelled; stock was not restored again.")
                    return redirect('order_cancellation')

                if order.status == MarketplaceOrder.STATUS_PENDING:
                    order.status = MarketplaceOrder.STATUS_CANCELLED
                    order.cancelled_at = timezone.now()
                    order.save(update_fields=['status', 'cancelled_at'])
                    messages.success(request, f"Pending order '{order_id}' cancelled. No stock was deducted.")
                    return redirect('order_cancellation')

                if order.status != MarketplaceOrder.STATUS_DISPATCHED:
                    messages.error(request, 'This order cannot be cancelled in its current status.')
                    return redirect('order_cancellation')
                if order.sku is None:
                    messages.error(request, 'This dispatched order has no linked SKU. Use Manual Restock.')
                    return redirect('order_cancellation')

                sku = order.sku
                batch = _restore_returned_stock(
                    sku,
                    order.quantity,
                    f"RST-{order.order_id}",
                    'Customer RTO/Return',
                )
                order.status = MarketplaceOrder.STATUS_CANCELLED
                order.cancelled_at = timezone.now()
                order.stock_restored = True
                order.save(update_fields=['status', 'cancelled_at', 'stock_restored'])
                sku.sync_stock_from_batches()

            messages.success(
                request,
                f"Order '{order_id}' cancelled. {order.quantity} unit(s) of {sku.sku} restored to batch {batch.batch_no}.",
            )
            return redirect('order_cancellation')

        if mode == 'manual':
            sku_id = request.POST.get('sku_id', '').strip()
            order_id = request.POST.get('manual_order_id', '').strip()
            quantity_raw = request.POST.get('quantity', '1').strip()
            if not sku_id.isdigit():
                messages.error(request, 'Choose a valid product.')
                return redirect('order_cancellation')
            try:
                quantity = int(quantity_raw)
                if quantity < 1:
                    raise ValueError
            except ValueError:
                messages.error(request, 'Please enter a valid quantity.')
                return redirect('order_cancellation')
            if len(order_id) > 100:
                messages.error(request, 'Order ID must be 100 characters or fewer.')
                return redirect('order_cancellation')

            sku = get_object_or_404(JewelrySKU, pk=sku_id)
            with transaction.atomic():
                if order_id and (
                    ManualStockReturn.objects.filter(order_id=order_id).exists()
                    or StockBatch.objects.filter(batch_no=f'RTO-{order_id}').exists()
                ):
                    messages.error(request, f"A return batch for Order ID '{order_id}' already exists.")
                    return redirect('order_cancellation')
                if order_id:
                    _, created = ManualStockReturn.objects.get_or_create(
                        order_id=order_id,
                        defaults={'sku': sku, 'quantity': quantity},
                    )
                    if not created:
                        messages.error(request, f"A return for Order ID '{order_id}' has already been processed.")
                        return redirect('order_cancellation')
                batch = _restore_returned_stock(
                    sku,
                    quantity,
                    f"RST-MANUAL-{timezone.now():%Y%m%d%H%M%S%f}",
                    'Manual Return/RTO',
                )
                sku.sync_stock_from_batches()

            messages.success(
                request,
                f'Manual return processed: {quantity} unit(s) added back to {sku.sku} in batch {batch.batch_no}.',
            )
            return redirect('order_cancellation')

        messages.error(request, 'Choose a valid return action.')
        return redirect('order_cancellation')

    products = JewelrySKU.objects.all().order_by('-created_at')
    cancelled_orders = (
        MarketplaceOrder.objects.filter(status=MarketplaceOrder.STATUS_CANCELLED)
        .select_related('sku')
        .order_by('-cancelled_at', '-created_at')[:15]
    )
    return render(request, 'sku_manager/cancellation.html', {
        'active_page': 'cancellation',
        'products': products,
        'platforms': Platform.objects.order_by('name'),
        'cancelled_orders': cancelled_orders,
    })


def _restore_returned_stock(sku, quantity, fallback_batch_no, supplier_name):
    batch = (
        sku.batches.filter(status='active').order_by('-created_at', '-pk').first()
        or sku.batches.order_by('-created_at', '-pk').first()
    )
    if batch:
        batch.quantity += quantity
        batch.status = 'active'
        batch.finished_at = None
        batch.save(update_fields=['quantity', 'status', 'finished_at', 'updated_at'])
        return batch

    return StockBatch.objects.create(
        sku=sku,
        batch_no=fallback_batch_no,
        quantity=quantity,
        purchase_price=sku.purchase_price,
        wholesale_price=sku.wholesale_price,
        selling_price=sku.selling_price,
        storage_box=sku.storage_box,
        section_name=sku.section_name,
        supplier_name=supplier_name,
        status='active',
    )


def parse_marketplace_order_ids(value):
    order_ids = []
    for part in value.split(','):
        try:
            order_id = int(part)
        except (TypeError, ValueError):
            continue
        if order_id > 0 and order_id not in order_ids:
            order_ids.append(order_id)
    return order_ids[:100]


@login_required(login_url='login')
def pdf_order_batch(request):
    raw_ids = request.POST.get('ids', '') if request.method == 'POST' else request.GET.get('ids', '')
    order_ids = parse_marketplace_order_ids(raw_ids)
    if not order_ids:
        messages.error(request, 'No scanned orders were selected for review.')
        return redirect('pdf_order_scan')

    raw_scan_id = request.POST.get('scan_id', '') if request.method == 'POST' else request.GET.get('scan_id', '')
    scan = None
    if raw_scan_id.isdigit():
        scan = MarketplaceLabelScan.objects.filter(
            pk=raw_scan_id,
            orders__pk__in=order_ids,
        ).distinct().first()

    if request.method == 'POST':
        dispatched_count = 0
        failed_count = 0
        available_ids = set(
            MarketplaceOrder.objects.filter(pk__in=order_ids).values_list('pk', flat=True)
        )
        for order_id in order_ids:
            if request.POST.get(f'include_{order_id}') != 'on':
                continue
            if order_id not in available_ids:
                failed_count += 1
                messages.error(request, f'Order {order_id}: scanned order was not found.')
                continue
            try:
                order, _ = dispatch_marketplace_order(
                    order_id=order_id,
                    sku_id=request.POST.get(f'sku_{order_id}'),
                    customer_name=request.POST.get(f'name_{order_id}', ''),
                    shipping_address=request.POST.get(f'address_{order_id}', ''),
                    save_seller_sku=request.POST.get(f'save_sku_{order_id}') == 'on',
                )
                dispatched_count += 1
            except MarketplaceOrderDispatchError as error:
                failed_count += 1
                messages.error(request, f"Order {order_id}: {error}")
            except IntegrityError:
                failed_count += 1
                messages.error(
                    request,
                    f"Order {order_id}: seller SKU is already mapped to another product "
                    f"on that marketplace.",
                )

        if dispatched_count:
            messages.success(request, f'{dispatched_count} order(s) dispatched successfully.')
        if not dispatched_count and not failed_count:
            messages.info(request, 'No orders were selected for dispatch.')
        query_ids = ','.join(str(order_id) for order_id in order_ids)
        query = f"{reverse('pdf_order_batch')}?ids={query_ids}"
        if scan:
            query += f'&scan_id={scan.pk}'
        return redirect(query)

    orders = list(
        MarketplaceOrder.objects.filter(pk__in=order_ids).select_related('platform', 'sku')
    )
    order_by_id = {order.pk: order for order in orders}
    orders = [order_by_id[order_id] for order_id in order_ids if order_id in order_by_id]
    if not orders:
        messages.error(request, 'No scanned orders were found.')
        return redirect('pdf_order_scan')

    rows = []
    for order in orders:
        matched = (
            order.sku
            if order.status == MarketplaceOrder.STATUS_DISPATCHED
            else find_product_for_seller_sku(order.platform, order.seller_sku)
        )
        rows.append({'order': order, 'matched': matched})

    return render(request, 'sku_manager/pdf_order_batch.html', {
        'active_page': 'pdf_order_scan',
        'rows': rows,
        'products': JewelrySKU.objects.filter(is_listed=True).order_by('name', 'sku'),
        'pending_count': sum(
            order.status == MarketplaceOrder.STATUS_PENDING for order in orders
        ),
        'ids': ','.join(str(order.pk) for order in orders),
        'skipped_pages': request.GET.get('skipped', ''),
        **marketplace_pdf_stamper_context(scan),
    })


@login_required(login_url='login')
def pdf_order_scan_file(request, pk):
    scan = get_object_or_404(MarketplaceLabelScan, pk=pk)
    return FileResponse(
        scan.pdf_file.open('rb'),
        as_attachment=False,
        filename=scan.pdf_file.name.rsplit('/', 1)[-1],
        content_type='application/pdf',
    )


@login_required(login_url='login')
def pdf_order_stamped_file(request, pk):
    scan = get_object_or_404(MarketplaceLabelScan, pk=pk)
    settings_obj = AppSettings.get_settings()
    include_logo = request.GET.get('logo') == '1' and bool(settings_obj.label_logo)
    include_store_qr = (
        request.GET.get('qr') == '1' and bool(settings_obj.meesho_store_url)
    )
    if not include_logo and not include_store_qr:
        return redirect('pdf_order_scan_file', pk=scan.pk)

    stamped_pdf = stamp_marketplace_pdf(scan, include_logo, include_store_qr)
    response = FileResponse(
        stamped_pdf,
        as_attachment=False,
        filename=f"stamped-{scan.pdf_file.name.rsplit('/', 1)[-1]}",
        content_type='application/pdf',
    )
    response['Cache-Control'] = 'private, no-store'
    return response


@login_required(login_url='login')
def pdf_order_review(request, pk):
    if request.method == 'POST':
        try:
            order, item = dispatch_marketplace_order(
                order_id=pk,
                sku_id=request.POST.get('sku_id'),
                customer_name=request.POST.get('customer_name', ''),
                shipping_address=request.POST.get('shipping_address', ''),
                save_seller_sku=request.POST.get('save_seller_sku') == 'on',
            )
        except (IntegrityError, MarketplaceOrderDispatchError) as error:
            if isinstance(error, IntegrityError):
                message = (
                    'That seller SKU is already mapped to another product on this '
                    'marketplace. Review the mapping before dispatching this order.'
                )
            else:
                message = str(error)
            messages.error(
                request,
                message,
            )
            return redirect('pdf_order_review', pk=pk)

        messages.success(
            request,
            f"Order {order.order_id} dispatched: {order.quantity} unit(s) of {item.sku}.",
        )
        return redirect('pdf_order_review', pk=order.pk)

    order = get_object_or_404(
        MarketplaceOrder.objects.select_related('platform', 'sku'),
        pk=pk,
    )
    scan = order.label_scans.first()
    products = JewelrySKU.objects.filter(is_listed=True).order_by('name', 'sku')
    if order.status == MarketplaceOrder.STATUS_PENDING:
        matched_product = find_product_for_seller_sku(order.platform, order.seller_sku)
    else:
        matched_product = order.sku
    return render(request, 'sku_manager/pdf_order_review.html', {
        'order': order,
        'products': products,
        'matched_product': matched_product,
        **marketplace_pdf_stamper_context(scan),
        'active_page': 'pdf_order_scan',
    })





@login_required(login_url='login')
def catalog_generator(request):
    """Generates a printable Catalog filtered by Category"""
    selected_category = request.GET.get('category', '').strip()

    # Fetch all unique categories from the database
    categories = JewelrySKU.objects.values_list('category', flat=True).distinct().order_by('category')

    # Fetch ALL products
    products = JewelrySKU.objects.all().order_by('-created_at')

    if selected_category:
        products = products.filter(category=selected_category)

    settings_obj = AppSettings.get_settings()

    return render(request, 'sku_manager/catalog.html', {
        'active_page': 'catalog',
        'categories': categories,
        'selected_category': selected_category,
        'products': products,
        'currency': settings_obj.currency_symbol,
        'brand_name': settings_obj.brand_name,
        'settings': settings_obj,  # Settings object passed for Logo
    })


@login_required(login_url='login')
def app_settings_view(request):
    settings_obj = AppSettings.get_settings()
    if request.method == 'POST':
        settings_obj.brand_name = request.POST.get('brand_name', '').strip() or 'TATKAL PICK'
        settings_obj.tagline = request.POST.get('tagline', '').strip()
        settings_obj.currency_symbol = request.POST.get('currency_symbol', '₹').strip()
        threshold = request.POST.get('low_stock_threshold', '5').strip()
        settings_obj.low_stock_threshold = int(threshold) if threshold.isdigit() else 5
        settings_obj.support_contact = request.POST.get('support_contact', '').strip()
        store_urls = {
            'meesho_store_url': request.POST.get('meesho_store_url', '').strip(),
            'amazon_store_url': request.POST.get('amazon_store_url', '').strip(),
            'flipkart_store_url': request.POST.get('flipkart_store_url', '').strip(),
        }
        validate_url = URLValidator(schemes=['http', 'https'])
        invalid_urls = []
        for field, value in store_urls.items():
            if value:
                try:
                    validate_url(value)
                except ValidationError:
                    invalid_urls.append(field.replace('_', ' ').replace(' store url', '').title())
            setattr(settings_obj, field, value)

        if invalid_urls:
            messages.error(
                request,
                'Enter a valid http:// or https:// URL for: ' + ', '.join(invalid_urls) + '.',
            )
            return render(request, 'sku_manager/settings.html', {
                'settings': settings_obj,
                'active_page': 'settings',
            })

        uploaded_logo = request.FILES.get('label_logo')
        if uploaded_logo:
            settings_obj.label_logo = uploaded_logo
        elif request.POST.get('remove_label_logo') and settings_obj.label_logo:
            settings_obj.label_logo.delete(save=False)
        settings_obj.save()
        messages.success(request, "Settings updated successfully!")
        return redirect('app_settings')

    return render(request, 'sku_manager/settings.html', {
        'settings': settings_obj,
        'active_page': 'settings',
    })